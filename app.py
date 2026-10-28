"""
HW3 Part B - RAG over your own course document with Python + MongoDB Atlas Vector Search.

ONE file, two ways to run it:

  Streamlit UI (recommended for screenshots):
      python -m streamlit run app.py

  Command line:
      python app.py --status                 DB / collection / documents / vector-index status  (B1, B2)
      python app.py --ingest FILE            extract, chunk, embed and store a document         (B2)
      python app.py --ask "question"         one RAG question with sources + timestamp          (B3, B4)
      python app.py --eval [questions.json]  retrieval evaluation (Hit@k, MRR)                  (B5)
      python app.py                          interactive question loop
      python app.py --seed                   (optional) the 4 sample policy docs from the starter kit
      python app.py --reset                  delete everything in the collection

"""
import argparse
import json
import os
import sys
import tempfile
from datetime import datetime
from urllib.parse import urlparse

from dotenv import load_dotenv
from openai import OpenAI
from pymongo import MongoClient
from sentence_transformers import SentenceTransformer

load_dotenv(override=True)

MONGODB_URI = os.getenv("MONGODB_URI", "mongodb://localhost:27017/")
DB_NAME = os.getenv("MONGODB_DB", "rag_demo")
COLLECTION_NAME = os.getenv("MONGODB_COLLECTION", "documents")
VECTOR_INDEX_NAME = os.getenv("VECTOR_INDEX_NAME", "vector_index")
EMBEDDING_MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

LLM_API_KEY = os.getenv("GROQ_API_KEY")
LLM_BASE_URL = os.getenv("LLM_BASE_URL", "https://api.groq.com/openai/v1")
CHAT_MODEL = os.getenv("CHAT_MODEL", "llama-3.3-70b-versatile")
# If CHAT_MODEL is retired / not available to your key (HTTP 404), the app asks Groq which models
# your key can use and picks the first of these that exists (else any chat model).
PREFERRED_MODELS = ["llama-3.3-70b-versatile", "openai/gpt-oss-120b", "openai/gpt-oss-20b",
                    "llama-3.1-8b-instant", "meta-llama/llama-4-scout-17b-16e-instruct", "qwen/qwen3-32b"]
NON_CHAT_WORDS = ("whisper", "guard", "tts", "embed", "orpheus", "safeguard")
_model_cache = {}
# 0.0 = off. If > 0, retrieved chunks scoring below this are dropped before the LLM is called.
MIN_SCORE = float(os.getenv("MIN_SCORE", "0.0"))

CHUNK_SIZE = 800      # characters per chunk
CHUNK_OVERLAP = 150   # characters shared between neighbouring chunks
TOP_K = 4             # chunks retrieved per question
EVAL_K = 5            # depth used by the retrieval evaluation
DEFAULT_EVAL_FILE = "eval_questions.json"

SAMPLE_DOCS = [
    {"title": "Refund Policy",
     "text": "Customers may request a refund within 30 days of purchase. The product must be in its original condition. Refunds are usually processed within 5 to 7 business days after approval."},
    {"title": "Shipping Policy",
     "text": "Standard shipping usually takes 3 to 5 business days. Express shipping usually takes 1 to 2 business days. Customers receive a tracking number after an order ships."},
    {"title": "Customer Support",
     "text": "Customer support is available Monday through Friday, 9 AM to 5 PM Central Time. Customers can contact support by emailing help@example.com. Include the order number when asking about an order."},
    {"title": "Warranty",
     "text": "Products include a one-year limited warranty covering manufacturing defects. The warranty does not cover accidental damage, misuse, or normal wear and tear."},
]


# ================================================================ helpers

def stamp():
    return datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")


def banner(title):
    print("=" * 76)
    print(f"{title}   |   {stamp()}")
    print("=" * 76)


def connect():
    client = MongoClient(MONGODB_URI, serverSelectionTimeoutMS=5000)
    client.admin.command("ping")
    return client, client[DB_NAME][COLLECTION_NAME]


def make_llm():
    return OpenAI(api_key=LLM_API_KEY, base_url=LLM_BASE_URL) if LLM_API_KEY else None


def label(item):
    page = item.get("page_number")
    name = item.get("source_file") or item.get("title", "?")
    return f"{name}, page {page}" if page else name


# ================================================================ status (B1 / B2)

def get_status(collection):
    """Setup info without secrets (no username, password or full URI)."""
    status = {
        "host": urlparse(MONGODB_URI).hostname,
        "db": DB_NAME,
        "collection": COLLECTION_NAME,
        "total_chunks": collection.count_documents({}),
        "files": [{"source_file": g["_id"], "chunks": g["chunks"]} for g in collection.aggregate([
            {"$group": {"_id": "$source_file", "chunks": {"$sum": 1}}},
            {"$sort": {"_id": 1}},
        ])],
    }
    first = collection.find_one({"embedding": {"$exists": True}})
    status["embedding_length"] = len(first["embedding"]) if first else None
    try:
        status["indexes"] = [{"name": ix.get("name"), "type": ix.get("type"),
                              "status": ix.get("status"), "queryable": ix.get("queryable")}
                             for ix in collection.list_search_indexes()]
        status["index_error"] = None
    except Exception as exc:
        status["indexes"] = []
        status["index_error"] = str(exc)
    return status


def show_status(collection):
    s = get_status(collection)
    banner("MongoDB status")
    print(f"Cluster host      : {s['host']}")
    print(f"Database          : {s['db']}")
    print(f"Collection        : {s['collection']}")
    print(f"Total chunks      : {s['total_chunks']}")
    for f in s["files"]:
        print(f"  - {f['source_file']}: {f['chunks']} chunks")
    print(f"Embedding length  : {s['embedding_length'] or '(none yet - ingest a file)'}")
    if s["index_error"]:
        print(f"Could not list search indexes: {s['index_error']}")
    if not s["indexes"]:
        print("Search indexes    : none found")
    for ix in s["indexes"]:
        print(f"Search index      : name={ix['name']} type={ix['type']} "
              f"status={ix['status']} queryable={ix['queryable']}")


# ================================================================ ingestion (B2)

def extract_pages(path):
    """Return a list of (page_number or None, text). PDF pages / PPTX slides keep their number."""
    ext = os.path.splitext(path)[1].lower()
    pages = []
    if ext == ".pdf":
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(path).pages, start=1):
            pages.append((i, page.extract_text() or ""))
    elif ext == ".pptx":
        from pptx import Presentation
        for i, slide in enumerate(Presentation(path).slides, start=1):
            parts = [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
            pages.append((i, "\n".join(parts)))
    elif ext == ".docx":
        from docx import Document
        pages.append((None, "\n".join(p.text for p in Document(path).paragraphs)))
    elif ext in (".txt", ".md"):
        with open(path, encoding="utf-8", errors="ignore") as f:
            pages.append((None, f.read()))
    else:
        raise ValueError(f"Unsupported file type: {ext} (use PDF, PPTX, DOCX, TXT or MD)")
    return pages


def chunk_text(text, size=CHUNK_SIZE, overlap=CHUNK_OVERLAP):
    """Split text into overlapping chunks, preferring to cut at a space."""
    text = " ".join(text.split())
    chunks, start = [], 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            cut = text.rfind(" ", start + size // 2, end)
            if cut != -1:
                end = cut
        chunk = text[start:end].strip()
        if len(chunk) >= 20:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def ingest_file(collection, embedder, path):
    """Extract -> chunk -> embed -> store. Re-ingesting a file replaces its old chunks.
    Returns a small summary dict."""
    filename = os.path.basename(path)
    pages = extract_pages(path)
    rows = []
    for page_number, text in pages:
        for idx, chunk in enumerate(chunk_text(text)):
            rows.append({
                "source_file": filename,
                "page_number": page_number,
                "chunk_index": idx,
                "title": f"{filename} (p.{page_number})" if page_number else filename,
                "text": chunk,
                "ingested_at": stamp(),
            })
    if not rows:
        raise ValueError("No extractable text found (scanned/image-only files need OCR).")
    vectors = embedder.encode([r["text"] for r in rows], normalize_embeddings=True,
                              batch_size=32, show_progress_bar=True).tolist()
    for row, vec in zip(rows, vectors):
        row["embedding"] = vec
    removed = collection.delete_many({"source_file": filename}).deleted_count
    collection.insert_many(rows)
    return {"filename": filename, "pages": len(pages), "chunks": len(rows),
            "replaced": removed, "sample": rows[0]["text"][:200], "time": stamp()}


def seed_documents(collection, embedder):
    """Optional: the starter kit's 4 sample policy documents."""
    banner("Seeding sample documents")
    vectors = embedder.encode([d["text"] for d in SAMPLE_DOCS], normalize_embeddings=True).tolist()
    collection.delete_many({"source_file": "sample_policies"})
    rows = [{**d, "source_file": "sample_policies", "page_number": None,
             "chunk_index": 0, "embedding": v} for d, v in zip(SAMPLE_DOCS, vectors)]
    collection.insert_many(rows)
    print(f"Loaded {len(rows)} sample documents.")


def reset_collection(collection):
    banner("Reset")
    n = collection.delete_many({}).deleted_count
    print(f"Deleted {n} documents from {DB_NAME}.{COLLECTION_NAME}.")


# ================================================================ retrieval + generation (B3 / B4)

def retrieve(collection, embedder, question, top_k=TOP_K):
    query_vector = embedder.encode([question], normalize_embeddings=True)[0].tolist()
    pipeline = [
        {"$vectorSearch": {
            "index": VECTOR_INDEX_NAME,
            "path": "embedding",
            "queryVector": query_vector,
            "numCandidates": max(100, top_k * 10),
            "limit": top_k,
        }},
        {"$project": {
            "_id": 0, "title": 1, "source_file": 1, "page_number": 1,
            "chunk_index": 1, "text": 1,
            "score": {"$meta": "vectorSearchScore"},
        }},
    ]
    return list(collection.aggregate(pipeline))


def available_models(llm):
    """Chat models this API key can actually use (asks the provider's /models endpoint)."""
    ids = [m.id for m in llm.models.list().data]
    return sorted(i for i in ids if not any(w in i.lower() for w in NON_CHAT_WORDS))


def chat_complete(llm, messages):
    """Call the LLM; on a 404 (model retired / no access) discover a working model and retry.
    Returns (text, model_name)."""
    model = _model_cache.get("name") or CHAT_MODEL
    try:
        resp = llm.chat.completions.create(model=model, temperature=0, messages=messages)
    except Exception as exc:
        if "404" not in str(exc) and "does not exist" not in str(exc):
            raise
        ids = available_models(llm)
        if not ids:
            raise RuntimeError("Your API key has no chat models available. Check console.groq.com "
                               "(Settings > Limits / model permissions).") from exc
        model = next((p for p in PREFERRED_MODELS if p in ids), ids[0])
        resp = llm.chat.completions.create(model=model, temperature=0, messages=messages)
    _model_cache["name"] = model
    return resp.choices[0].message.content, model


def rag_answer(collection, embedder, llm, question):
    """Retrieve -> augment -> generate. Returns a dict with answer, retrieved chunks, timestamp."""
    retrieved = retrieve(collection, embedder, question)
    used = [m for m in retrieved if m.get("score", 0) >= MIN_SCORE] if MIN_SCORE > 0 else retrieved
    model = None
    if not used:
        answer = ("I don't know based on the available documents "
                  f"(no chunk scored above MIN_SCORE={MIN_SCORE}).")
    elif llm is None:
        answer = "(GROQ_API_KEY is not set - showing retrieval only.)"
    else:
        context = "\n\n".join(f"Source: {label(m)}\nContent: {m['text']}" for m in used)
        answer, model = chat_complete(llm, [
            {"role": "system", "content": (
                "Answer the user's question using only the supplied context. "
                "If the context does not contain the answer, say you don't know "
                "based on the available documents. Cite the source file and page you used.")},
            {"role": "user", "content": f"Context:\n{context}\n\nQuestion: {question}"},
        ])
    return {"question": question, "answer": answer, "retrieved": retrieved,
            "used": used, "time": stamp(), "model": model}


def answer_question(collection, embedder, llm, question):
    banner(f"Question: {question}")
    result = rag_answer(collection, embedder, llm, question)
    print("Retrieved sources:")
    for item in result["retrieved"]:
        snippet = item["text"][:110].replace("\n", " ")
        print(f"- {label(item)} (score {item.get('score', 0):.3f}): {snippet}...")
    print("\nAnswer:")
    print(result["answer"])
    if result.get("model"):
        print(f"\n(model: {result['model']})")


# ================================================================ retrieval evaluation (B5)

def is_relevant(chunk, spec, default_source):
    """A retrieved chunk is relevant if it is from the right file and on an expected page
    or contains the expected keyword."""
    source = spec.get("source_file") or default_source
    if source and chunk.get("source_file") != source:
        return False
    pages = spec.get("expected_pages") or []
    keyword = (spec.get("expected_keyword") or "").lower()
    if pages and chunk.get("page_number") in pages:
        return True
    if keyword and keyword in chunk.get("text", "").lower():
        return True
    return False


def run_evaluation(collection, embedder, cfg):
    default_source = cfg.get("source_file")
    rows, ranks = [], []
    for spec in cfg["questions"]:
        results = retrieve(collection, embedder, spec["question"], top_k=EVAL_K)
        rank = next((i for i, c in enumerate(results, 1) if is_relevant(c, spec, default_source)), None)
        ranks.append(rank)
        rows.append({"question": spec["question"], "rank": rank,
                     "top1_score": round(results[0]["score"], 3) if results else 0,
                     "top1_source": label(results[0]) if results else "-"})
    n = len(ranks)
    hit = lambda cut: sum(1 for r in ranks if r is not None and r <= cut) / n if n else 0
    unsupported = []
    for q in cfg.get("unsupported", []):
        top = retrieve(collection, embedder, q, top_k=1)
        if top:
            unsupported.append({"question": q, "top1_source": label(top[0]),
                                "top1_score": round(top[0]["score"], 3)})
    return {"rows": rows, "n": n, "hit1": hit(1), "hit3": hit(3), "hit5": hit(5),
            "mrr": sum(1 / r for r in ranks if r) / n if n else 0,
            "unsupported": unsupported, "time": stamp()}


def evaluate(collection, embedder, path):
    banner(f"Retrieval evaluation ({path})")
    if not os.path.exists(path):
        print(f"{path} not found. Copy eval_questions.example.json to {path} and fill in "
              "your own questions (see the file for the format).")
        return
    with open(path, encoding="utf-8") as f:
        cfg = json.load(f)
    res = run_evaluation(collection, embedder, cfg)
    print(f"{'Rank':<5} {'Top-1 score':<12} Question")
    for r in res["rows"]:
        print(f"{str(r['rank']):<5} {r['top1_score']:<12.3f} {r['question']}")
        print(f"      top-1 was: {r['top1_source']}")
    print("-" * 76)
    print(f"Questions: {res['n']}   (rank = position of first relevant chunk; None = not in top {EVAL_K})")
    print(f"Hit@1: {res['hit1']:.2f}   Hit@3: {res['hit3']:.2f}   Hit@5: {res['hit5']:.2f}   MRR: {res['mrr']:.2f}")
    if res["unsupported"]:
        print("\nUnsupported questions (expect LOWER top-1 scores than the supported ones):")
        for u in res["unsupported"]:
            print(f"- {u['question']}  -> top-1 {u['top1_source']} (score {u['top1_score']:.3f})")


# ================================================================ Streamlit UI

def run_streamlit():
    import streamlit as st

    st.set_page_config(page_title="Course RAG - MongoDB", page_icon="📚", layout="wide")
    st.title("📚 Course Document Q&A")
    st.caption(f"RAG with MongoDB Atlas Vector Search  |  System time: {stamp()}")

    # --- connections (kept per browser session)
    if "collection" not in st.session_state:
        try:
            st.session_state.client, st.session_state.collection = connect()
        except Exception as exc:
            st.error(f"Could not connect to MongoDB: {exc}")
            st.info("Check MONGODB_URI in .env, the database user's password, and Network Access in Atlas.")
            st.stop()
    if "embedder" not in st.session_state:
        with st.spinner("Loading embedding model (first time takes a minute)..."):
            st.session_state.embedder = SentenceTransformer(EMBEDDING_MODEL_NAME)
    collection, embedder = st.session_state.collection, st.session_state.embedder
    llm = make_llm()
    if llm is None:
        st.warning("GROQ_API_KEY is not set in .env - answers will show retrieval only.")

    # --- sidebar: status + upload
    with st.sidebar:
        st.header("MongoDB status")
        status = get_status(collection)
        st.write(f"**Host:** {status['host']}")
        st.write(f"**Database / collection:** {status['db']}.{status['collection']}")
        st.write(f"**Total chunks:** {status['total_chunks']}")
        st.write(f"**Embedding length:** {status['embedding_length'] or '-'}")
        for f in status["files"]:
            st.write(f"- {f['source_file']}: {f['chunks']} chunks")
        if status["index_error"]:
            st.error(status["index_error"])
        elif not status["indexes"]:
            st.warning("No vector search index found yet. Create it in Atlas (see README).")
        for ix in status["indexes"]:
            st.write(f"**Index:** {ix['name']} - {ix['status']} (queryable: {ix['queryable']})")

        st.divider()
        st.header("Add a document")
        uploaded = st.file_uploader("PDF, PPTX, DOCX, TXT, MD", type=["pdf", "pptx", "docx", "txt", "md"])
        if uploaded and st.button("Ingest document", type="primary"):
            path = os.path.join(tempfile.mkdtemp(), uploaded.name)  # keeps the original file name
            with open(path, "wb") as f:
                f.write(uploaded.getvalue())
            try:
                with st.spinner("Extracting, chunking, embedding, storing..."):
                    info = ingest_file(collection, embedder, path)
                st.success(f"{info['filename']}: {info['pages']} page(s) -> {info['chunks']} chunks "
                           f"stored at {info['time']}")
                st.rerun()
            except Exception as exc:
                st.error(str(exc))

    page = st.sidebar.radio("Page", ["Ask questions", "Retrieval evaluation"])

    def render_sources(retrieved):
        for i, m in enumerate(retrieved, 1):
            st.markdown(f"**{i}. {label(m)}** - similarity `{m['score']:.3f}`")
            st.caption(m["text"][:500] + ("..." if len(m["text"]) > 500 else ""))

    # --- page 1: chat
    if page == "Ask questions":
        st.session_state.setdefault("history", [])
        for turn in st.session_state.history:
            with st.chat_message("user"):
                st.write(turn["question"])
            with st.chat_message("assistant"):
                st.write(turn["answer"])
                st.caption(f"Answered at {turn['time']}" + (f"  |  model: {turn['model']}" if turn.get("model") else ""))
                with st.expander("Retrieved sources"):
                    render_sources(turn["retrieved"])
        question = st.chat_input("Ask a question about your document")
        if question:
            with st.chat_message("user"):
                st.write(question)
            with st.chat_message("assistant"):
                try:
                    with st.spinner("Searching MongoDB and generating an answer..."):
                        turn = rag_answer(collection, embedder, llm, question)
                except Exception as exc:
                    st.error(f"Could not answer: {exc}")
                    st.info("Confirm the Atlas Vector Search index exists and is Active.")
                    st.stop()
                st.write(turn["answer"])
                st.caption(f"Answered at {turn['time']}" + (f"  |  model: {turn['model']}" if turn.get("model") else ""))
                with st.expander("Retrieved sources", expanded=True):
                    render_sources(turn["retrieved"])
            st.session_state.history.append(turn)

    # --- page 2: evaluation
    else:
        st.subheader("Retrieval evaluation")
        st.write("Does the right chunk come back near the top? Each question lists the page(s) or "
                 "keyword where the answer lives (see `eval_questions.example.json`).")
        cfg_file = st.file_uploader("Questions file (JSON)", type=["json"])
        cfg = None
        if cfg_file:
            cfg = json.load(cfg_file)
        elif os.path.exists(DEFAULT_EVAL_FILE):
            st.info(f"Using {DEFAULT_EVAL_FILE} from the project folder.")
            with open(DEFAULT_EVAL_FILE, encoding="utf-8") as f:
                cfg = json.load(f)
        else:
            st.warning(f"No {DEFAULT_EVAL_FILE} found. Upload one, or copy the example file and edit it.")
        if cfg and st.button("Run evaluation", type="primary"):
            res = run_evaluation(collection, embedder, cfg)
            st.caption(f"Run at {res['time']}")
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("Hit@1", f"{res['hit1']:.2f}")
            c2.metric("Hit@3", f"{res['hit3']:.2f}")
            c3.metric("Hit@5", f"{res['hit5']:.2f}")
            c4.metric("MRR", f"{res['mrr']:.2f}")
            st.table([{"Question": r["question"], "Rank of right chunk": r["rank"],
                       "Top-1 score": r["top1_score"], "Top-1 source": r["top1_source"]}
                      for r in res["rows"]])
            if res["unsupported"]:
                st.markdown("**Unsupported questions** (expect lower top-1 scores)")
                st.table([{"Question": u["question"], "Top-1 source": u["top1_source"],
                           "Top-1 score": u["top1_score"]} for u in res["unsupported"]])


# ================================================================ command line

def main():
    parser = argparse.ArgumentParser(description="MongoDB RAG over your course document")
    parser.add_argument("--status", action="store_true", help="show DB and index status")
    parser.add_argument("--ingest", metavar="FILE", help="ingest a PDF/PPTX/DOCX/TXT/MD file")
    parser.add_argument("--ask", metavar="QUESTION", help="ask one question and exit")
    parser.add_argument("--eval", nargs="?", const=DEFAULT_EVAL_FILE, metavar="FILE",
                        help=f"run retrieval evaluation (default {DEFAULT_EVAL_FILE})")
    parser.add_argument("--models", action="store_true", help="list chat models your GROQ_API_KEY can use")
    parser.add_argument("--seed", action="store_true", help="load the 4 sample policy docs")
    parser.add_argument("--reset", action="store_true", help="delete all documents in the collection")
    args = parser.parse_args()

    if args.models:
        if not LLM_API_KEY:
            print("Set GROQ_API_KEY in your .env file first.")
            sys.exit(1)
        banner("Chat models available to your API key")
        for m in available_models(make_llm()):
            print("-", m)
        return

    try:
        client, collection = connect()
    except Exception as exc:
        print(f"Could not connect to MongoDB: {exc}")
        print("Check MONGODB_URI and confirm your Atlas cluster / IP allowlist.")
        sys.exit(1)

    try:
        if args.status:
            show_status(collection)
            return
        if args.reset:
            reset_collection(collection)
            return

        print("Loading embedding model...")
        embedder = SentenceTransformer(EMBEDDING_MODEL_NAME)

        if args.ingest:
            banner(f"Ingesting {args.ingest}")
            info = ingest_file(collection, embedder, args.ingest)
            print(f"Extracted {info['pages']} page(s)/slide(s) -> {info['chunks']} chunks "
                  f"(size {CHUNK_SIZE}, overlap {CHUNK_OVERLAP})")
            print(f"Stored in {DB_NAME}.{COLLECTION_NAME} (replaced {info['replaced']} old chunks).")
            print("Sample chunk:", info["sample"].replace("\n", " "), "...")
            print("\nNext: make sure the vector index is Active, then run: python app.py --status")
            return
        if args.seed:
            seed_documents(collection, embedder)
            return
        if args.eval is not None:
            evaluate(collection, embedder, args.eval)
            return

        if not LLM_API_KEY:
            print("Set GROQ_API_KEY in your .env file before asking questions.")
            sys.exit(1)
        llm = make_llm()

        if args.ask:
            answer_question(collection, embedder, llm, args.ask)
            return

        print(f"Connected to {DB_NAME}.{COLLECTION_NAME}")
        print("Ask a question about your document. Type 'exit' to quit.")
        while True:
            question = input("\nQuestion: ").strip()
            if question.lower() in {"exit", "quit"}:
                break
            if question:
                try:
                    answer_question(collection, embedder, llm, question)
                except Exception as exc:
                    print(f"Could not answer question: {exc}")
                    print("Confirm that the Atlas Vector Search index exists and is ACTIVE.")
    finally:
        client.close()


def _running_in_streamlit():
    try:
        from streamlit.runtime.scriptrunner import get_script_run_ctx
        try:
            return get_script_run_ctx(suppress_warning=True) is not None
        except TypeError:
            return get_script_run_ctx() is not None
    except Exception:
        return False


if __name__ == "__main__":
    if _running_in_streamlit():
        run_streamlit()
    else:
        main()
